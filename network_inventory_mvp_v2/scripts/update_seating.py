#!/usr/bin/env python3
"""Update the Agent Name column of an Excel seating map from movement events.

Only movement events are applied. Static seat fields (Seat ID, Desk Port, Hostname,
Switch Port and Switch Serial) are never changed.

Safety behavior:
* destination seat must resolve uniquely by current switch + current interface;
* a destination occupied by a different user is never overwritten;
* only users with PORT_MOVED/SWITCH_MOVED statuses are processed;
* old occurrences of the same Everest ID are cleared only after destination
  validation succeeds;
* a cell comment is added to the new Agent Name cell to identify the move.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Tuple

try:
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
except ImportError:
    print(
        "ERROR: openpyxl is required for seating-map updates. "
        "Add openpyxl>=3.1.0 to the AWX Execution Environment.",
        file=sys.stderr,
    )
    raise SystemExit(3)

SWITCH_RE = re.compile(r"((?:ASW|DSW|CSW)\d+)", re.I)
IF_RE = re.compile(r"^(ge|xe|et|fe|mge)-?(\d+)/(\d+)/(\d+)(?:\.\d+)?$", re.I)
MOVE_PREFIXES = ("PORT_MOVED", "SWITCH_MOVED")


def norm_text(value) -> str:
    return "" if value is None else str(value).strip()


def norm_switch(value) -> str:
    text = norm_text(value)
    match = SWITCH_RE.search(text)
    return match.group(1).lower() if match else text.lower()


def norm_interface(value) -> str:
    text = norm_text(value).replace(" ", "")
    match = IF_RE.match(text)
    if not match:
        return text.lower()
    return f"{match.group(1).lower()}-{match.group(2)}/{match.group(3)}/{match.group(4)}"


def norm_header(value) -> str:
    return re.sub(r"\s+", " ", norm_text(value)).casefold()


def read_events(path: Path) -> List[Dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def find_seating_sheet(workbook, requested_sheet: str = ""):
    if requested_sheet:
        if requested_sheet not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet '{requested_sheet}' was not found. Available sheets: "
                + ", ".join(workbook.sheetnames)
            )
        candidates = [workbook[requested_sheet]]
    else:
        candidates = list(workbook.worksheets)

    required = {
        "seat id (do not edit)",
        "desk port",
        "agent name",
        "switch port",
        "switch serial",
    }

    for sheet in candidates:
        for row_idx in range(1, min(sheet.max_row, 20) + 1):
            header_map: Dict[str, int] = {}
            for col_idx in range(1, sheet.max_column + 1):
                key = norm_header(sheet.cell(row=row_idx, column=col_idx).value)
                if key:
                    header_map[key] = col_idx
            if required.issubset(header_map):
                return sheet, row_idx, header_map

    raise ValueError(
        "Could not find a seating table with headers: "
        "Seat ID (do not edit), Desk Port, Agent Name, Switch Port, Switch Serial."
    )


def build_indexes(sheet, header_row: int, columns: Dict[str, int]):
    location_index: Dict[Tuple[str, str], List[int]] = {}
    agent_index: Dict[str, List[int]] = {}

    for row_idx in range(header_row + 1, sheet.max_row + 1):
        switch = norm_switch(sheet.cell(row_idx, columns["switch serial"]).value)
        interface = norm_interface(sheet.cell(row_idx, columns["switch port"]).value)
        agent = norm_text(sheet.cell(row_idx, columns["agent name"]).value).casefold()

        if switch and interface:
            location_index.setdefault((switch, interface), []).append(row_idx)
        if agent:
            agent_index.setdefault(agent, []).append(row_idx)

    return location_index, agent_index


def seat_snapshot(sheet, row_idx: int, columns: Dict[str, int]) -> Dict[str, str]:
    return {
        "seat_id": norm_text(sheet.cell(row_idx, columns["seat id (do not edit)"]).value),
        "desk_port": norm_text(sheet.cell(row_idx, columns["desk port"]).value),
        "agent_name": norm_text(sheet.cell(row_idx, columns["agent name"]).value),
        "switch_port": norm_text(sheet.cell(row_idx, columns["switch port"]).value),
        "switch_serial": norm_text(sheet.cell(row_idx, columns["switch serial"]).value),
    }


def is_move(status: str) -> bool:
    return any(status.startswith(prefix) for prefix in MOVE_PREFIXES)


def write_csv(path: Path, rows: List[Dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "everest_id",
        "movement_status",
        "update_status",
        "source_seat_id",
        "source_desk_port",
        "destination_seat_id",
        "destination_desk_port",
        "destination_switch",
        "destination_interface",
        "message",
    ]
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--events", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--report", required=True)
    parser.add_argument("--summary-json", required=True)
    parser.add_argument("--sheet", default="")
    parser.add_argument("--job-id", default="manual")
    parser.add_argument("--client", default="")
    parser.add_argument("--comment-author", default="AWX Network Inventory")
    args = parser.parse_args()

    source_path = Path(args.workbook)
    events_path = Path(args.events)
    output_path = Path(args.output)
    report_path = Path(args.report)
    summary_path = Path(args.summary_json)

    if not source_path.exists():
        print(f"ERROR: Seating workbook not found: {source_path}", file=sys.stderr)
        return 2
    if not events_path.exists():
        print(f"ERROR: Movement report not found: {events_path}", file=sys.stderr)
        return 2

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, output_path)

    workbook = load_workbook(output_path)
    try:
        sheet, header_row, columns = find_seating_sheet(workbook, args.sheet)
    except ValueError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    events = [row for row in read_events(events_path) if is_move(row.get("status", ""))]
    location_index, agent_index = build_indexes(sheet, header_row, columns)
    report: List[Dict[str, str]] = []
    timestamp = datetime.now(timezone.utc).isoformat()

    updated_count = 0
    conflict_count = 0
    destination_not_found_count = 0
    destination_ambiguous_count = 0

    for event in events:
        user = norm_text(event.get("everest_id"))
        user_key = user.casefold()
        movement_status = norm_text(event.get("status"))
        current_switch = norm_switch(event.get("current_switch"))
        current_interface = norm_interface(event.get("current_interface"))

        base_report = {
            "everest_id": user,
            "movement_status": movement_status,
            "update_status": "",
            "source_seat_id": "",
            "source_desk_port": "",
            "destination_seat_id": "",
            "destination_desk_port": "",
            "destination_switch": norm_text(event.get("current_switch")),
            "destination_interface": norm_text(event.get("current_interface")),
            "message": "",
        }

        destination_rows = location_index.get((current_switch, current_interface), [])
        if not destination_rows:
            destination_not_found_count += 1
            report.append(
                {
                    **base_report,
                    "update_status": "DESTINATION_NOT_FOUND",
                    "message": "Current switch/port was not found in the seating table.",
                }
            )
            continue
        if len(destination_rows) > 1:
            destination_ambiguous_count += 1
            report.append(
                {
                    **base_report,
                    "update_status": "DESTINATION_AMBIGUOUS",
                    "message": f"Current switch/port matches {len(destination_rows)} seating rows.",
                }
            )
            continue

        destination_row = destination_rows[0]
        destination = seat_snapshot(sheet, destination_row, columns)
        destination_agent_cell = sheet.cell(destination_row, columns["agent name"])
        destination_existing = norm_text(destination_agent_cell.value)

        source_rows = [row for row in agent_index.get(user_key, []) if row != destination_row]
        source = seat_snapshot(sheet, source_rows[0], columns) if source_rows else {
            "seat_id": "",
            "desk_port": "",
            "agent_name": "",
            "switch_port": norm_text(event.get("baseline_interface")),
            "switch_serial": norm_text(event.get("baseline_switch")),
        }

        base_report.update(
            {
                "source_seat_id": source.get("seat_id", ""),
                "source_desk_port": source.get("desk_port", ""),
                "destination_seat_id": destination.get("seat_id", ""),
                "destination_desk_port": destination.get("desk_port", ""),
            }
        )

        if destination_existing and destination_existing.casefold() != user_key:
            conflict_count += 1
            report.append(
                {
                    **base_report,
                    "update_status": "SEAT_CONFLICT",
                    "message": (
                        f"Destination seat is already assigned to '{destination_existing}'. "
                        "No cells were changed."
                    ),
                }
            )
            continue

        # Destination is safe. Clear any stale occurrences of the same user first.
        cleared_rows: List[str] = []
        for source_row in source_rows:
            snapshot = seat_snapshot(sheet, source_row, columns)
            sheet.cell(source_row, columns["agent name"]).value = None
            cleared_rows.append(snapshot.get("seat_id", "") or str(source_row))

        destination_agent_cell.value = user
        comment_lines = [
            "Movement detected automatically by AWX Network Inventory.",
            f"User: {user}",
            f"Status: {movement_status}",
            f"Detected UTC: {timestamp}",
            f"AWX Job ID: {args.job_id}",
        ]
        if args.client:
            comment_lines.append(f"Client: {args.client}")
        if source.get("seat_id"):
            comment_lines.append(
                f"Previous seat: {source['seat_id']} ({source.get('desk_port') or 'no desk port'})"
            )
        comment_lines.append(
            f"Previous network location: {norm_text(event.get('baseline_switch'))} "
            f"{norm_text(event.get('baseline_interface'))}"
        )
        comment_lines.append(
            f"Current seat: {destination.get('seat_id') or '?'} "
            f"({destination.get('desk_port') or 'no desk port'})"
        )
        comment_lines.append(
            f"Current network location: {destination.get('switch_serial')} "
            f"{destination.get('switch_port')}"
        )
        if cleared_rows:
            comment_lines.append("Cleared previous seat row(s): " + ", ".join(cleared_rows))

        destination_agent_cell.comment = Comment("\n".join(comment_lines), args.comment_author)
        updated_count += 1
        report.append(
            {
                **base_report,
                "update_status": "UPDATED",
                "message": (
                    f"Agent Name updated to {user}. "
                    f"Previous seat(s) cleared: {', '.join(cleared_rows) if cleared_rows else 'none found'}."
                ),
            }
        )

        # Keep indexes coherent for additional movement events in the same job.
        agent_index[user_key] = [destination_row]

    workbook.save(output_path)
    write_csv(report_path, report)

    summary = {
        "source_workbook": str(source_path),
        "output_workbook": str(output_path),
        "worksheet": sheet.title,
        "movement_events_received": len(events),
        "updated": updated_count,
        "seat_conflicts": conflict_count,
        "destination_not_found": destination_not_found_count,
        "destination_ambiguous": destination_ambiguous_count,
        "report": str(report_path),
    }
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")

    print(f"Seating worksheet: {sheet.title}")
    print(f"Movement events received: {len(events)}")
    print(f"Seats updated: {updated_count}")
    print(f"Seat conflicts: {conflict_count}")
    print(f"Destination not found: {destination_not_found_count}")
    print(f"Destination ambiguous: {destination_ambiguous_count}")
    for row in report:
        if row["update_status"] == "UPDATED":
            print(
                f"  UPDATED: {row['everest_id']} | "
                f"Seat {row['source_seat_id'] or '?'} -> {row['destination_seat_id'] or '?'} | "
                f"{row['destination_switch']}:{row['destination_interface']}"
            )
        else:
            print(
                f"  {row['update_status']}: {row['everest_id']} | {row['message']}"
            )
    print(f"Updated workbook: {output_path}")
    print(f"Seating update report: {report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
